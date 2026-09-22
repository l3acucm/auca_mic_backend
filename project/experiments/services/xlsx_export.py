"""XLSX export (BRD module 4 / 6.2). Single-participant exports are stored in
S3 (private, signed URL to download); the "all participants of one
experiment" export is generated on demand and streamed straight back —
nothing to store for a report nobody asked to keep."""
import io
from datetime import datetime, timezone

import structlog
from django.core.files.base import ContentFile
from openpyxl import Workbook

from project.experiments.services.scoring import CORRECT, ERROR, INCORRECT, SKIPPED

logger = structlog.get_logger(__name__)

_RESULT_HEADERS = ['Стимул', 'Время реакции (сек)', 'Кодирование', 'Распознанный текст']
_SUMMARY_HEADERS = [
    'ID участника', 'Дата и время попытки', 'Количество попыток выполненных',
    'Общее время попытки (сек)', 'Количество верных ответов',
    'Количество неверных ответов', 'Количество пропусков', 'Количество ошибок',
]


def aggregate_trials(trial_data: list[dict]) -> dict:
    counts = {CORRECT: 0, INCORRECT: 0, SKIPPED: 0, ERROR: 0}
    for trial in trial_data:
        counts[trial['code']] = counts.get(trial['code'], 0) + 1
    total_time = 0.0
    if trial_data:
        # t_конец_последней_попытки - t_первый_стимул (BRD 3.10). The end of
        # the last trial is its own stimulus timestamp + its reaction time —
        # works uniformly whether it ended in speech, a skip, or a timeout.
        first_ts = datetime.fromisoformat(trial_data[0]['timestamp_stimulus'].replace('Z', '+00:00'))
        last_trial = trial_data[-1]
        last_ts = datetime.fromisoformat(last_trial['timestamp_stimulus'].replace('Z', '+00:00'))
        total_time = (last_ts - first_ts).total_seconds() + last_trial['reaction_time_sec']
    return {
        'num_correct': counts[CORRECT], 'num_incorrect': counts[INCORRECT],
        'num_skipped': counts[SKIPPED], 'num_errors': counts[ERROR],
        'total_time_sec': round(total_time, 3),
    }


def build_result_workbook(participant_id: str, trial_data: list[dict], aggregates: dict, completed_at) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Результаты'
    ws.append(_RESULT_HEADERS)
    for trial in trial_data:
        ws.append([
            trial['stimulus_filename'], round(trial['reaction_time_sec'], 3),
            trial['code'], trial.get('recognized_text', ''),
        ])

    summary = wb.create_sheet('Сводка')
    summary.append(_SUMMARY_HEADERS)
    summary.append([
        participant_id, completed_at.isoformat(), len(trial_data),
        aggregates['total_time_sec'], aggregates['num_correct'],
        aggregates['num_incorrect'], aggregates['num_skipped'], aggregates['num_errors'],
    ])
    return wb


def result_filename(participant_id: str, experiment_id) -> str:
    stamp = datetime.now(timezone.utc).strftime('%Y-%m-%d_%H-%M-%S')
    return f'{participant_id}_{experiment_id}_{stamp}.xlsx'


def save_result_workbook(wb: Workbook, participant_id: str, filename: str) -> str:
    """Uploads to PrivateMediaStorage under /results/{participant_id}/ and
    returns the storage path (BRD 4.4)."""
    from project.storages import PrivateMediaStorage

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    storage = PrivateMediaStorage()
    return storage.save(f'results/{participant_id}/{filename}', ContentFile(buf.getvalue()))


def build_experiment_workbook(experiment, results) -> Workbook:
    """All completed results of one experiment: a combined summary sheet + a
    combined per-trial sheet (BRD 6.2)."""
    wb = Workbook()
    summary = wb.active
    summary.title = 'Сводка'
    summary.append(_SUMMARY_HEADERS)
    for result in results:
        summary.append([
            result.participant_id, result.completed_at.isoformat(), len(result.trial_data),
            result.total_time_sec, result.num_correct, result.num_incorrect,
            result.num_skipped, result.num_errors,
        ])

    trials = wb.create_sheet('Все ответы')
    trials.append(['ID участника'] + _RESULT_HEADERS)
    for result in results:
        for trial in result.trial_data:
            trials.append([
                result.participant_id, trial['stimulus_filename'],
                round(trial['reaction_time_sec'], 3), trial['code'],
                trial.get('recognized_text', ''),
            ])
    return wb
